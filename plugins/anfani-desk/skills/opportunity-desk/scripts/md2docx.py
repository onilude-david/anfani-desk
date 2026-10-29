"""Convert an answer sheet (Markdown) to a Word document the user can open, print or upload to Google Drive.

Usage: python md2docx.py <answers.md> <out.docx>
Handles headings, paragraphs, bold, italics, inline code, links (as text), bullet and numbered lists, tables,
blockquotes and horizontal rules. [TBC: ...] markers are highlighted yellow.
"""
import re, sys
try:
    from docx import Document
except ImportError:
    import sys
    sys.exit('error: md2docx needs python-docx. Install it with:  python -m pip install --user python-docx')
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[TBC[^\]]*\]|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*]))')

def add_runs(par, text, base_italic=False):
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', r'\1 (\2)', text)
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'
        elif part.startswith('[TBC'):
            r = par.add_run(part); r.bold = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if base_italic: r.italic = True

def convert(md, out):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    lines = md.splitlines(); i = 0; para = []
    def flush():
        if para:
            add_runs(doc.add_paragraph(), ' '.join(para)); para.clear()
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            flush(); i += 1; continue
        m = re.match(r'^(#{1,6})\s+(.*)$', ln)
        if m:
            flush(); h = doc.add_heading('', level=min(len(m.group(1)), 4)); add_runs(h, m.group(2)); i += 1; continue
        if re.match(r'^\s*(-{3,}|\*{3,})\s*$', ln):
            flush(); doc.add_paragraph('_' * 40); i += 1; continue
        if ln.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-+:?$', c) for c in cells if c): rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'
                for r, cells in enumerate(rows):
                    for c in range(ncol):
                        cell = t.cell(r, c); cell.text = ''
                        p = cell.paragraphs[0]; add_runs(p, cells[c] if c < len(cells) else '')
                        if r == 0:
                            for run in p.runs: run.bold = True
                doc.add_paragraph()
            continue
        if ln.startswith('>'):
            flush(); q = []
            while i < len(lines) and lines[i].startswith('>'):
                q.append(lines[i].lstrip('> ').rstrip()); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(24); add_runs(p, ' '.join(q)); continue
        m = re.match(r'^(\s*)([-*]|\d+[.)])\s+(.*)$', ln)
        if m:
            flush(); ordered = m.group(2)[0].isdigit()
            while i < len(lines):
                m2 = re.match(r'^(\s*)([-*]|\d+[.)])\s+(.*)$', lines[i])
                if m2 and (m2.group(2)[0].isdigit() == ordered):
                    item = m2.group(3); i += 1
                    while i < len(lines) and lines[i].startswith('  ') and lines[i].strip() and not re.match(r'^\s*([-*]|\d+[.)])\s+', lines[i]):
                        item += ' ' + lines[i].strip(); i += 1
                    add_runs(doc.add_paragraph(style='List Number' if ordered else 'List Bullet'), item)
                else: break
            continue
        para.append(ln.strip()); i += 1
    flush(); doc.save(out)

if __name__ == '__main__':
    src, dst = sys.argv[1:3]
    convert(open(src, encoding='utf-8').read(), dst); print(dst)
